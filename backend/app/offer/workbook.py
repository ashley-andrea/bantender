"""The full pricing as an Excel workbook with live formulas, for the estimator who wants to check every step.

Sheets: Overview · Parameters · Catalogue · Productivity · Positions · Build-up · Offer form · Formulas.
Every number the app computes is a formula here that points at its inputs: change a parameter in Excel and the
positions, the build-up, the offer price and the form follow. Values the app takes as given (catalogue net prices,
hours per unit, supplier quotes, the hygiene allocation) are plain values with their source named next to them.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from app.pricing.engine import HEIGHT_PREFIX, SITE_EFFICIENCY
from app.reference.models import ReferenceData

NAVY, GREEN, MINT, PAPER = "002439", "008163", "DDEFE5", "F7F5EF"
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")  # values the estimator may change
TOTAL_FILL = PatternFill("solid", fgColor=MINT)
THIN = Border(bottom=Side(style="thin", color="D9DEE2"))
CHF = '#,##0.00'
CHF0 = '#,##0'
PCT = '0.00" %"'


def _header(ws, cols: list[tuple[str, int]], row: int = 1) -> None:
    for i, (title, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = f"A{row + 1}"  # a string: ws.cell() would create an empty row and shift append()
    ws.row_dimensions[row].height = 30


def build_workbook(parsed: dict, pricing: dict, offer: dict, ref: ReferenceData) -> bytes:
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    s = pricing["settings"]
    rows = pricing["lines"]
    cost_rows = [r for r in rows if r["cost"]]

    # ---- Parameters ------------------------------------------------------------------------------
    wp = wb.active
    wp.title = "Parameters"
    _header(wp, [("Parameter", 34), ("Value", 14), ("Unit", 10), ("Group", 18), ("Source", 34)])
    cell: dict[str, str] = {}
    defaults = ref.parameters

    def param(key: str, value: float, unit: str, group: str, default: float | None) -> None:
        r = wp.max_row + 1
        wp.cell(r, 1, key)
        v = wp.cell(r, 2, value)
        v.fill = INPUT_FILL
        v.number_format = "0.00" if unit in ("×", "h") else CHF if unit.startswith("CHF") else "0.00"
        wp.cell(r, 3, unit)
        wp.cell(r, 4, group)
        changed = default is not None and abs(default - value) > 1e-9
        wp.cell(r, 5, f"estimator (company value {default:g})" if changed else "project_parameters.json")
        cell[key] = f"Parameters!$B${r}"

    for k, v in s["factors"].items():
        param(k, v, "×", "Productivity factor", defaults.productivity_factors.get(k))
    for k, v in s["allowances"].items():
        param(k, v, "%", "Allowance", getattr(defaults.allowances, k, None))
    for k, v in s["project_costs"].items():
        param(k, v, "h" if k.endswith("_hours") else "CHF/h" if "rate" in k else "CHF", "Project cost", getattr(defaults.project_costs, k, None))
    param("overhead_pct", s["overhead_pct"], "%", "Commercial", defaults.commercial.overhead_pct)
    param("margin_pct", s["margin_pct"], "%", "Commercial", defaults.commercial.target_margin_pct)
    param("risk_allowance_chf", s["risk_allowance_chf"], "CHF", "Commercial", 0.0)
    param("round_to_chf", s["round_to_chf"], "CHF", "Commercial", defaults.commercial.round_to_chf)
    flags = {"include_project_costs": s["include_project_costs"], "include_small_material": s["include_small_material"],
             "include_supplier_contingency": s["include_supplier_contingency"]}
    for k, v in flags.items():
        param(k, 1.0 if v else 0.0, "1 = yes", "Switch", 1.0)
    r = wp.max_row + 2
    wp.cell(r, 1, "Yellow cells are inputs: change them and every sheet follows.").font = Font(italic=True, color="5D6B75")
    r += 1
    wp.cell(r, 1, "Mounting height per system").font = Font(bold=True)
    for code, loc in s["locations"].items():
        r += 1
        wp.cell(r, 1, f"{code} {parsed['systems'][code]['name']}")
        wp.cell(r, 2, loc)

    # ---- Catalogue (rows used by this tender) ------------------------------------------------------
    wc = wb.create_sheet("Catalogue")
    _header(wc, [("Product", 14), ("Category", 18), ("Subcategory", 22), ("Description", 44), ("Unit", 7),
                 ("List CHF", 12), ("Discount %", 11), ("Net CHF", 12), ("Install code", 18), ("Source", 18)])
    used = sorted({r["decision"]["product_id"] for r in rows if r["decision"]["product_id"]})
    cat_row: dict[str, int] = {}
    for pid in used:
        item = ref.catalog.get(pid)
        if not item:
            continue
        wc.append([item.product_id, item.category, item.subcategory, item.description, item.unit, item.list_price_chf,
                   item.discount_pct, item.net_cost_chf, item.install_code, item.source_type])
        cat_row[pid] = wc.max_row
        for col, fmt in ((6, CHF), (8, CHF)):
            wc.cell(wc.max_row, col).number_format = fmt
    wc.cell(wc.max_row + 2, 1, "From catalog.csv (company reference data). Net = list × (1 − discount).").font = Font(italic=True, color="5D6B75")

    # ---- Productivity (rules used) -----------------------------------------------------------------
    wr = wb.create_sheet("Productivity")
    _header(wr, [("Install code", 20), ("Work type", 30), ("Unit", 7), ("Hours / unit", 12), ("Labour role", 18),
                 ("Rate CHF/h", 11), ("Project factors", 46)])
    codes = sorted({r["decision"]["install_code"] or (r["cost"] or {}).get("install_code") for r in rows} - {None})
    rule_row: dict[str, int] = {}
    for code in codes:
        rule = ref.rules.get(code)
        if not rule:
            continue
        wr.append([rule.install_code, rule.work_type, rule.unit, rule.base_hours_per_unit, rule.labour_role,
                   rule.default_hourly_rate_chf, ", ".join(rule.factor_keys)])
        rule_row[code] = wr.max_row
    wr.cell(wr.max_row + 2, 1, "From productivity.csv (company reference data).").font = Font(italic=True, color="5D6B75")

    # ---- Positions: one row per priced line, every step a formula -----------------------------------
    wl = wb.create_sheet("Positions")
    cols = [("Line", 18), ("Source", 7), ("System", 8), ("Section", 8), ("Page", 6), ("Position", 34), ("Position (tender)", 30),
            ("Status", 20), ("Product", 13), ("Quantity", 10), ("Unit", 6), ("Net unit CHF", 13), ("Price source", 16),
            ("Waste %", 9), ("Material CHF", 14), ("Install code", 18), ("Hours / unit", 11), ("Factor", 9),
            ("Hours", 10), ("Rate CHF/h", 10), ("Labour CHF", 13), ("Direct cost CHF", 15), ("Factor composition", 52)]
    _header(wl, cols)
    pos_row: dict[str, int] = {}
    for row in cost_rows:
        d, c = row["decision"], row["cost"]
        n = wl.max_row + 1
        pos_row[row["line_id"]] = n
        q = d["quantity"] or 0.0
        pid = d["product_id"]
        overridden = "unit_price" in d.get("overridden", [])
        if overridden or not pid or pid not in cat_row:
            net = c["net_unit_cost"]
        else:
            net = f"=Catalogue!$H${cat_row[pid]}"
        code = c.get("install_code") or d["install_code"]
        rule = ref.rules.get(code) if code else None
        hours_override = "hours_per_unit" in d.get("overridden", [])
        if rule is None or c["adjusted_hours"] == 0:
            base, factor, rate = 0, 1, 0
            composition = "no installation hours"
        else:
            base = c["base_hours"] if hours_override or code not in rule_row else f"=Productivity!$D${rule_row[code]}"
            parts, names = [], []
            for key in rule.factor_keys:
                loc = d.get("location", "default")
                if key.startswith(HEIGHT_PREFIX) and loc != "default" and f"{HEIGHT_PREFIX}{loc}" in s["factors"]:
                    key = f"{HEIGHT_PREFIX}{loc}"
                parts.append(f"1/{cell[key]}" if key == SITE_EFFICIENCY else cell[key])
                names.append(f"1/{key}" if key == SITE_EFFICIENCY else key)
            factor = "=" + "*".join(parts) if parts else 1
            rate = f"=Productivity!$F${rule_row[code]}" if code in rule_row else c["hourly_rate"]
            composition = " × ".join(names)
        wl.append([row["line_id"], row["source"].upper(), row["system"], row["section"], row["page"], row["title_en"], row["title_de"],
                   d["status"], pid or "", q, d["unit"] or "", net, c["price_source"], f"={cell['material_waste_pct']}",
                   f"=J{n}*L{n}*(1+N{n}/100)", code or "", base, factor, f"=J{n}*Q{n}*R{n}", rate, f"=S{n}*T{n}", f"=O{n}+U{n}",
                   composition + (" · hours per unit set by the estimator" if hours_override else "")
                   + (" · net price from a supplier quote" if overridden else "")])
        for col, fmt in ((12, CHF), (14, "0.00"), (15, CHF), (17, "0.00"), (18, "0.0000"), (19, "0.00"), (21, CHF), (22, CHF)):
            wl.cell(n, col).number_format = fmt
        if overridden:
            wl.cell(n, 12).fill = INPUT_FILL
        if hours_override:
            wl.cell(n, 17).fill = INPUT_FILL
    last = wl.max_row
    total = last + 1
    wl.cell(total, 6, "Total").font = Font(bold=True)
    for col in (15, 21, 22):
        letter = get_column_letter(col)
        t = wl.cell(total, col, f"=SUM({letter}2:{letter}{last})")
        t.font, t.fill, t.number_format = Font(bold=True), TOTAL_FILL, CHF
    excluded = [r for r in rows if not r["cost"] and r["source"] == "lv"]
    r = total + 2
    wl.cell(r, 1, "Positions without a cost line").font = Font(bold=True)
    for row in excluded:
        r += 1
        wl.cell(r, 1, row["line_id"])
        wl.cell(r, 6, row["title_en"])
        wl.cell(r, 8, row["decision"]["status"])

    # ---- Build-up: positions → offer price ---------------------------------------------------------
    wbu = wb.create_sheet("Build-up")
    _header(wbu, [("Step", 34), ("CHF", 16), ("Formula", 60)])
    rng = lambda col: f"Positions!${col}$2:${col}${last}"  # noqa: E731
    steps: list[tuple[str, str, str]] = [
        ("Material of all positions", f"=SUM({rng('O')})", "Σ material (Positions)"),
        ("of which supplier-package material", f"=SUMIF({rng('M')},\"supplier_package\",{rng('O')})", "Σ material where price source = supplier_package"),
        ("of which catalogue material", "=B2-B3", "material − supplier-package material"),
        ("Installation of all positions", f"=SUM({rng('U')})", "Σ labour (Positions)"),
        ("Positions direct cost", "=B2+B5", "material + installation"),
        ("Transport / equipment", f"=B6*{cell['transport_equipment_pct']}/100", "transport % × positions direct cost"),
        ("Small material", f"=B4*{cell['small_material_pct']}/100*{cell['include_small_material']}", "small material % × catalogue material"),
        ("Supplier contingency", f"=B3*{cell['supplier_package_contingency_pct']}/100*{cell['include_supplier_contingency']}",
         "contingency % × supplier-package material"),
        ("Engineering", f"={cell['engineering_hours']}*{cell['engineering_rate_chf_per_h']}*{cell['include_project_costs']}", "hours × rate"),
        ("Commissioning", f"={cell['commissioning_hours']}*{cell['commissioning_rate_chf_per_h']}*{cell['include_project_costs']}", "hours × rate"),
        ("Hygiene inspection", f"={cell['hygiene_inspection_chf']}*{cell['include_project_costs']}", "fixed amount"),
        ("Documentation", f"={cell['documentation_chf']}*{cell['include_project_costs']}", "fixed amount"),
        ("Direct cost", "=B6+B7+B8+B9+B10+B11+B12+B13", "positions + transport + small material + contingency + project costs"),
        ("Risk / opportunities", f"={cell['risk_allowance_chf']}", "estimator input"),
        ("Overhead", f"=(B14+B15)*{cell['overhead_pct']}/100", "overhead % × (direct cost + risk)"),
        ("Cost baseline", "=B14+B15+B16", "direct cost + risk + overhead"),
        ("Price before rounding", f"=B17/(1-{cell['margin_pct']}/100)", "baseline ÷ (1 − margin %)"),
        ("Offer price", f"=CEILING(B18,{cell['round_to_chf']})", "rounded up to the rounding step"),
        ("Surcharge factor m", "=B19/B14", "offer price ÷ direct cost (SIA 118 Art. 39: one factor on every unit price)"),
    ]
    for label, formula, text in steps:
        wbu.append([label, formula, text])
        wbu.cell(wbu.max_row, 2).number_format = CHF if label != "Surcharge factor m" else "0.000000"
    for rr in (14, 17, 19):
        for cc in (1, 2):
            wbu.cell(rr, cc).font = Font(bold=True)
            wbu.cell(rr, cc).fill = TOTAL_FILL
    wbu.cell(19, 2).number_format = CHF0

    # ---- Offer form: what is written into the tender ----------------------------------------------
    wo = wb.create_sheet("Offer form")
    _header(wo, [("Line", 18), ("System", 8), ("Section", 8), ("Position", 34), ("Quantity", 10), ("Material CHF", 14),
                 ("Allowance CHF", 14), ("Installation on the form CHF", 16), ("Before factor CHF", 15), ("Factor m", 10),
                 ("Unit price CHF", 14), ("Amount CHF", 14), ("Note", 46)])
    by_id = {r["line_id"]: r for r in rows}
    allocated: dict[str, list[int]] = {}  # system → form rows that carry an allocated project cost (hygiene)
    for line_id, ln in offer["lines"].items():
        row = by_id.get(line_id)
        if not row or ln["kind"] == "empty":
            continue
        n = wo.max_row + 1
        if ln["kind"] == "included":
            wo.append([line_id, row["system"], row["section"], row["title_en"], None, None, None, None, None, None, "incl.", "incl.",
                       "included in another position"])
            continue
        p = pos_row.get(line_id)
        if p is None:  # e.g. the hygiene inspection line: an allocation of a project cost, not a position cost
            allocated.setdefault(row["system"], []).append(n)
            wo.append([line_id, row["system"], row["section"], row["title_en"], ln["quantity"], None, None, None, None, None,
                       ln["unit_price"], ln["amount"], "allocated share of a project cost (see Build-up)"])
        else:
            pct = (f"IF(Positions!$M${p}=\"supplier_package\",{cell['supplier_package_contingency_pct']}*{cell['include_supplier_contingency']},"
                   f"{cell['small_material_pct']}*{cell['include_small_material']})")
            labour = f"=Positions!$U${p}" if ln.get("includes_labour") else 0
            wo.append([line_id, row["system"], row["section"], row["title_en"], f"=Positions!$J${p}", f"=Positions!$O${p}",
                       f"=F{n}*{pct}/100", labour, f"=F{n}+G{n}+H{n}", "='Build-up'!$B$20",
                       f"=IF(E{n}>0,ROUND(ROUND(I{n}*J{n},2)/E{n},2),ROUND(I{n}*J{n},2))", f"=IF(E{n}>0,ROUND(K{n}*E{n},2),K{n})",
                       "installation included (section F / labour only)" if ln.get("includes_labour")
                       else "installation goes to section E of this system"])
        for col, fmt in ((6, CHF), (7, CHF), (8, CHF), (9, CHF), (10, "0.000000"), (11, CHF), (12, CHF)):
            wo.cell(n, col).number_format = fmt
    last_o = wo.max_row
    r = last_o + 2
    wo.cell(r, 1, "Per system (as written on the form)").font = Font(bold=True)
    r += 1
    head = ["System", "", "", "", "A–D, F positions", "Installation + transport", "Services", "Hygiene", "CAD material list", "", "", "Total"]
    for i, h in enumerate(head, start=1):
        if h:
            c = wo.cell(r, i, h)
            c.font = Font(bold=True)
    first_sys = r + 1
    for code, sysd in offer["systems"].items():
        r += 1
        wo.cell(r, 1, code)
        wo.cell(r, 4, parsed["systems"][code]["name"])
        minus = "".join(f"-L{a}" for a in allocated.get(code, []))  # the hygiene line is its own column
        wo.cell(r, 5, f"=SUMIF($B$2:$B${last_o},\"{code}\",$L$2:$L${last_o}){minus}")
        wo.cell(r, 6, sysd["montage"])
        wo.cell(r, 7, sysd["services"])
        wo.cell(r, 8, sysd["hygiene"])
        wo.cell(r, 9, sysd["cad_material"])
        wo.cell(r, 12, f"=E{r}+F{r}+G{r}+H{r}+I{r}")
        for col in (5, 6, 7, 8, 9, 12):
            wo.cell(r, col).number_format = CHF
    last_sys = r
    r += 1
    wo.cell(r, 4, "Gross total").font = Font(bold=True)
    g = wo.cell(r, 12, f"=SUM(L{first_sys}:L{last_sys})")
    g.font, g.fill, g.number_format = Font(bold=True), TOTAL_FILL, CHF
    gross = f"L{r}"
    c = offer["conditions"]
    cond = [("Discount %", c["discount_pct"], None), ("Discount", None, f"=ROUND({gross}*L{{p}}/100,2)"),
            ("Subtotal 1", None, f"={gross}-L{{d}}"), ("Cash discount %", c["cash_discount_pct"], None),
            ("Cash discount", None, "=ROUND(L{s1}*L{p2}/100,2)"), ("Subtotal 2", None, "=L{s1}-L{cd}"),
            ("Deductions %", c["deductions_pct"], None), ("Deductions", None, "=ROUND(L{s2}*L{p3}/100,2)"),
            ("Subtotal 3", None, "=L{s2}-L{de}"), ("VAT %", c["vat_pct"], None), ("VAT", None, "=ROUND(L{s3}*L{p4}/100,2)"),
            ("Net incl. VAT", None, "=L{s3}+L{v}")]
    names = {}
    keys = ["p", "d", "s1", "p2", "cd", "s2", "p3", "de", "s3", "p4", "v", "net"]
    for (label, value, formula), key in zip(cond, keys):
        r += 1
        names[key] = r
        wo.cell(r, 4, label)
        cl = wo.cell(r, 12, value if formula is None else formula.format(**names))
        cl.number_format = "0.00" if "%" in label else CHF
        if formula is None:
            cl.fill = INPUT_FILL
    wo.cell(r, 4).font = wo.cell(r, 12).font = Font(bold=True)
    wo.cell(r, 12).fill = TOTAL_FILL
    wo.cell(r + 2, 1, "Installation, services and hygiene are spread over the systems by their share of the direct cost; "
                      "they are shown as values (see Build-up for their totals).").font = Font(italic=True, color="5D6B75")

    # ---- Formulas (plain language) -----------------------------------------------------------------
    wf = wb.create_sheet("Formulas")
    _header(wf, [("Step", 30), ("Formula", 90)])
    for step, formula in [
        ("Net unit price", "catalogue net price (list × (1 − discount)), or the supplier quote the estimator entered"),
        ("Material", "quantity × net unit price × (1 + material waste %)"),
        ("Installation hours", "quantity × hours per unit × productivity factors (site efficiency as 1 / x; mounting height by system)"),
        ("Installation", "installation hours × hourly rate of the labour role"),
        ("Direct cost of a position", "material + installation"),
        ("Transport / equipment", "transport % × Σ positions direct cost"),
        ("Small material", "small material % × catalogue material"),
        ("Supplier contingency", "contingency % × supplier-package material"),
        ("Project costs", "engineering hours × rate + commissioning hours × rate + hygiene inspection + documentation"),
        ("Direct cost", "positions + transport + small material + contingency + project costs"),
        ("Overhead", "overhead % × (direct cost + risk)"),
        ("Cost baseline", "direct cost + risk + overhead"),
        ("Offer price", "baseline ÷ (1 − margin %), rounded up to the rounding step"),
        ("Surcharge factor m", "offer price ÷ direct cost: one factor on every cost component (SIA 118 Art. 39)"),
        ("Position on the form", "(material + its allowance [+ installation for section F and labour-only lines]) × m, rounded to centimes"),
        ("Unit price on the form", "position amount ÷ quantity, rounded to centimes; amount = unit price × quantity"),
        ("Section E", "installation of sections A–D + transport share, services and hygiene, per system"),
        ("Conditions (p.2)", "gross − discount − cash discount − deductions, + VAT = net"),
    ]:
        wf.append([step, formula])

    # ---- Overview (first sheet) --------------------------------------------------------------------
    wv = wb.create_sheet("Overview", 0)
    wv.column_dimensions["A"].width = 34
    wv.column_dimensions["B"].width = 60
    wv["A1"] = f"{parsed.get('project', {}).get('name') or 'Tender'} · BKP {parsed.get('project', {}).get('bkp') or ''}"
    wv["A1"].font = Font(bold=True, size=14, color=NAVY)
    wv["A3"], wv["B3"] = "Offer price (CHF)", "='Build-up'!B19"
    wv["A4"], wv["B4"] = "Direct cost (CHF)", "='Build-up'!B14"
    wv["A5"], wv["B5"] = "Margin", f"={cell['margin_pct']}/100"
    wv["A6"], wv["B6"] = "Net incl. VAT on the form (CHF)", f"='Offer form'!L{names['net']}"
    for ref_cell, fmt in (("B3", CHF0), ("B4", CHF0), ("B5", "0.0%"), ("B6", CHF)):
        wv[ref_cell].number_format = fmt
        wv[ref_cell].alignment = Alignment(horizontal="left")
    wv["A3"].font = wv["B3"].font = Font(bold=True)
    rr = 8
    for name, what in [("Parameters", "every coefficient (yellow = input; change it and all sheets follow)"),
                       ("Catalogue", "the catalogue rows this tender uses (catalog.csv)"),
                       ("Productivity", "the installation rules this tender uses (productivity.csv)"),
                       ("Positions", "each position: material, hours, installation, direct cost, as formulas"),
                       ("Build-up", "from the positions to the offer price"),
                       ("Offer form", "unit prices and amounts written into the tender, per system, conditions"),
                       ("Formulas", "every step in words")]:
        wv.cell(rr, 1, name).font = Font(bold=True)
        wv.cell(rr, 2, what)
        rr += 1
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title in ("Positions", "Offer form")
        for row in ws.iter_rows(min_row=2):
            for cl in row:
                if cl.border != THIN and ws.title != "Overview":
                    cl.border = THIN

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
